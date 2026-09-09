"""Portfolio research workspace. No order-placement code; all drafts are local.

Provider adapters produce validated, ISIN-keyed snapshots without overwriting
legacy holdings. Broker lookups only use the verified paper Gateway.
"""
import asyncio
import csv
import html
import io
import json
import math
import re
import shutil
import tempfile
import threading
import time
import urllib.parse
import urllib.request
import uuid
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

HERE = Path(__file__).resolve().parent
DATA = HERE / '.workspace'
LOCK = threading.RLock()
BROKER_LOCK = threading.Lock()
JOBS = {}
# Official product identifiers, verified against the issuer's ISIN field.
ISHARES = {
    'IE00BD1F4L37': '285209', 'IE00B4K48X80': '251861',
    'IE00BMG6Z448': '315592', 'IE00BQT3WG13': '273192',
    'IE00BMX0DF60': '314925', 'IE00B6R52143': '251707',
    'IE00B1TXK627': '251913',
}
ISHARES_SLUGS = {
    '285209': 'ishares-edge-msci-usa-quality-factor-ucits-etf-fund',
    '251861': 'ishares-msci-europe-ucits-etf-acc-fund',
    '315592': 'ishares-msci-em-ex-china-ucits-etf',
    '273192': 'ishares-msci-china-a-ucits-etf',
    '314925': 'ishares-us-medical-devices-ucits-etf-fund',
    '251707': 'ishares-agribusiness-ucits-etf',
    '251913': 'ishares-global-water-ucits-etf',
}
GLOBALX = {'IE0003Z9E2Y3': 'copx', 'IE00BLCHJB90': 'botz'}
SSGA = {'IE00BSPLC413': 'https://www.ssga.com/uk/en_gb/institutional/etfs/'
        'state-street-spdr-msci-usa-small-cap-value-weighted-ucits-etf-zprv-gy'}
ALLOWED_HOSTS = ('ishares.com', 'blackrock.com', 'etf.dws.com',
                 'etf.dws.com.cn', 'globalxetfs.eu', 'ssga.com',
                 'invesco.com', 'vaneck.com', 'lgim.com', 'fundcentres.lgim.com',
                 'ubs.com', 'ftglobalportfolios.com', 'firsttrust.com', 'ftportfolios.com')
MAX_DOWNLOAD = 25 * 1024 * 1024


def now():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')


def atomic_json(path, data):
    path.parent.mkdir(exist_ok=True)
    with tempfile.NamedTemporaryFile(mode='w', dir=path.parent, delete=False) as f:
        json.dump(data, f, allow_nan=False, indent=2)
        name = f.name
    Path(name).replace(path)


def read_json(path, default):
    try:
        return json.loads(path.read_text())
    except (OSError, ValueError):
        return default


def catalog():
    entries = []
    identities = read_json(DATA/'identities.json', {})
    for ticker, s in json.loads((HERE / 'manual.json').read_text())['sleeves'].items():
        entries.append(dict(ticker=ticker, name=s['name'], isin=s['isin'],
            uid='isin:' + s['isin'], currency=s['currency'], exchange=s['exchange'],
            weight=s['target'], ter=s.get('ter'), conid=None,
            identity_source='Operating manual', instrument_type='ETF',
            override='CASH:EUR' if ticker == 'XEON' else 'GOLD:GOLD' if ticker == 'SGLD' else '',
            source_url='', provider_page=ishares_page(s['isin'])))
        identity = identities.get(s['isin']+':'+s['currency'])
        if identity:
            entries[-1].update(identity)
    return entries


def ishares_page(isin):
    if isin in SSGA:
        return SSGA[isin]
    if isin in GLOBALX:
        return f'https://globalxetfs.eu/funds/{GLOBALX[isin]}/'
    pid = ISHARES.get(isin)
    return f'https://www.ishares.com/ch/professionals/en/products/{pid}/{ISHARES_SLUGS[pid]}' if pid else ''


def valid_isin(s):
    if not re.fullmatch(r'[A-Z]{2}[A-Z0-9]{9}[0-9]', s):
        return False
    digits = ''.join(str(ord(c)-55) if c.isalpha() else c for c in s)
    return sum((int(d)*2)//10+(int(d)*2)%10 if i%2 else int(d)
               for i,d in enumerate(reversed(digits))) % 10 == 0


def validate_rows(rows, require_total=False):
    if not isinstance(rows, list) or not 1 <= len(rows) <= 60:
        raise ValueError('Add between 1 and 60 investments.')
    clean, seen = [], set()
    for r in rows:
        isin = str(r.get('isin', '')).strip().upper()
        if not valid_isin(isin):
            raise ValueError('Choose an investment with a valid ISIN before continuing.')
        if isin in seen:
            raise ValueError('The same ISIN appears twice. Combine its allocations into one row.')
        seen.add(isin)
        weight = float(r.get('weight', 0))
        if not math.isfinite(weight) or not 0 < weight <= 100:
            raise ValueError('Each allocation must be greater than 0 and at most 100%.')
        ticker = str(r.get('ticker', '')).strip().upper()
        if not re.fullmatch(r'[A-Z0-9][A-Z0-9. _-]{0,23}', ticker):
            raise ValueError('Invalid ticker. Choose an investment from search.')
        entry = {k: str(r.get(k, ''))[:500] for k in
                 ('name','currency','exchange','identity_source','instrument_type','source_url','provider_page')}
        entry.update(isin=isin, uid='isin:'+isin, ticker=ticker, weight=weight,
                     conid=int(r['conid']) if r.get('conid') else None)
        manual = next((s for s in catalog() if s['isin']==isin), None)
        entry['override'] = manual['override'] if manual else ''
        entry['ter'] = manual['ter'] if manual else None
        clean.append(entry)
    if require_total and abs(sum(r['weight'] for r in clean)-100) > .01:
        raise ValueError('Allocations must add up to 100% to run the X-Ray.')
    return clean


def with_broker(fn):
    with BROKER_LOCK:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        from ib_async import IB
        ib = IB()
        try:
            ib.connect('127.0.0.1', 4002, clientId=72, timeout=5, readonly=True)
            accounts = ib.managedAccounts()
            if not accounts or not all(a.startswith('DU') for a in accounts):
                raise ValueError('Instrument lookup requires the paper Gateway on port 4002.')
            ib.RequestTimeout = 9
            return fn(ib)
        finally:
            ib.disconnect()
            loop.close()
            asyncio.set_event_loop(None)


def detail_record(d):
    c = d.contract
    isin = next((s.value for s in d.secIdList if s.tag == 'ISIN'), '')
    base = next((r for r in catalog() if r['isin'] == isin), {})
    return {**base, 'ticker': c.localSymbol or c.symbol, 'name': d.longName or c.symbol,
        'isin': isin, 'uid': 'isin:'+isin if isin else 'ibkr:'+str(c.conId),
        'conid': c.conId, 'currency': c.currency, 'exchange': c.primaryExchange or c.exchange,
        'identity_source': 'IBKR verified', 'instrument_type': d.stockType or 'STK',
        'weight': base.get('weight', 0), 'provider_page': ishares_page(isin),
        'source_url': '', 'override': base.get('override','')}


def search(query, online=False):
    query = str(query).strip()[:80]
    if len(query) < 2:
        return {'ok':True, 'candidates':[], 'note':'Enter at least two characters.'}
    q = query.upper()
    matches = [r for r in catalog() if q in (r['ticker']+' '+r['name']+' '+r['isin']).upper()]
    note = 'Saved fund identities from your operating manual. Use broker search to compare listings.'
    if online:
        from ib_async import Contract
        def lookup(ib):
            if valid_isin(q):
                return [detail_record(d) for d in ib.reqContractDetails(
                    Contract(secType='STK', secIdType='ISIN', secId=q, exchange='SMART'))]
            # IB limits matching-symbol requests to one per second. This lock
            # spans connection/search, and the sleep enforces that interval.
            time.sleep(1.05)
            return [dict(ticker=c.contract.symbol, name=c.contract.description or c.contract.symbol,
                         conid=c.contract.conId, isin='', currency=c.contract.currency,
                         exchange=c.contract.primaryExchange, identity_source='IBKR search',
                         instrument_type=c.contract.secType)
                    for c in ib.reqMatchingSymbols(query) if c.contract.secType == 'STK']
        try:
            found = with_broker(lookup)
            matches += [r for r in found if r.get('conid') not in {x.get('conid') for x in matches if x.get('conid')}]
            note = 'Choose the exact fund and listing. Different exchanges and currencies can share a ticker.'
        except Exception as e:
            note = f'Broker search unavailable ({type(e).__name__}). Check paper Gateway port 4002. Saved funds remain available.'
    return {'ok':True, 'candidates':matches[:30], 'note':note}


def resolve(p):
    from ib_async import Contract
    conid = int(p.get('conid') or 0)
    isin = str(p.get('isin', '')).upper()
    currency = str(p.get('currency','')).upper()
    if not conid and not valid_isin(isin):
        raise ValueError('A valid ISIN or broker contract ID is required.')
    def lookup(ib):
        c = Contract(conId=conid, exchange='SMART') if conid else Contract(
            secType='STK', secIdType='ISIN', secId=isin, exchange='SMART', currency=currency)
        return [detail_record(d) for d in ib.reqContractDetails(c)]
    records = with_broker(lookup)
    return {'ok':True, 'candidates':records,
            'note':'Identity verified with IBKR. Choose the listing you intend to use.'}


def allowed_url(url):
    p = urllib.parse.urlsplit(url)
    host = (p.hostname or '').lower()
    if p.scheme != 'https' or p.username or p.password or p.port not in (None,443) or not any(
            host == h or host.endswith('.'+h) for h in ALLOWED_HOSTS):
        raise ValueError('Use an HTTPS download or product page from a supported ETF provider.')
    return url


class ProviderRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        allowed_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def fetch(url):
    allowed_url(url)
    opener = urllib.request.build_opener(ProviderRedirect)
    req = urllib.request.Request(url, headers={'User-Agent':'Mozilla/5.0', 'Accept':'*/*'})
    with opener.open(req, timeout=18) as r:
        b = r.read(MAX_DOWNLOAD + 1)
        if len(b) > MAX_DOWNLOAD:
            raise ValueError('Provider file exceeds the 25 MB limit.')
        return b, r.geturl()


def xml_holdings(data, isin):
    ns = {'s':'urn:schemas-microsoft-com:office:spreadsheet'}
    root = ET.fromstring(data)
    # Fund downloads contain share-class ISIN in Key Facts, not just holding ISINs.
    sheets = {}
    for sh in root.findall('s:Worksheet', ns):
        rows = []
        for row in sh.findall('s:Table/s:Row', ns):
            cells = []
            for c in row.findall('s:Cell', ns):
                index = int(c.get('{'+ns['s']+'}Index',len(cells)+1))
                cells += [''] * max(0,index-len(cells)-1)
                d = c.find('s:Data',ns)
                cells.append(''.join(d.itertext()) if d is not None else '')
            rows.append(cells)
        sheets[sh.get('{'+ns['s']+'}Name')] = rows
    facts = sheets.get('Key Facts',[])
    actual = next((r[1] for r in facts if len(r)>1 and r[0].strip().upper()=='ISIN'), '')
    if actual != isin:
        raise ValueError(f'Provider identity mismatch: expected {isin}, received {actual or "no ISIN"}.')
    rows = sheets.get('Holdings',[])
    if not rows:
        raise ValueError('The provider workbook has no Holdings sheet.')
    output = io.StringIO()
    csv.writer(output).writerows(rows)
    asof = next((r[1] for r in rows[:5] if len(r)>1 and r[0].lower()=='as of'), None)
    return output.getvalue().encode(), asof


def snapshot_path(isin):
    return DATA / (isin + '.csv')


def source_meta(row):
    isin = row['isin']
    if row.get('override'):
        return {'status':'classified', 'label':'Cash / gold classification', 'as_of':None,
                'note':'Economic exposure from the operating manual; no equity look-through.'}
    m = read_json(DATA / (isin+'.json'), {})
    if snapshot_path(isin).exists() and m:
        return {**m, 'status':'downloaded', 'label':'Provider download'}
    # Legacy filenames are only trusted when bound to the manual's exact ISIN.
    manual = next((r for r in catalog() if r['isin']==isin), None)
    if manual:
        for ext in ('.csv','.xlsx','.tsv'):
            p = HERE/'holdings'/(manual['ticker']+ext)
            if p.exists():
                text = p.read_bytes()[:500].decode('utf-8-sig','replace') if ext!='.xlsx' else ''
                m = re.search(r'(\d{1,2}/[A-Za-z]{3,4}/\d{4})',text)
                return {'status':'local','label':'Local file', 'file':p.name,
                        'as_of':m.group(1) if m else None,
                        'note':'Previously imported file. Provider identity and freshness have not been reverified.'}
    return {'status':'missing','label':'Needs holdings','as_of':None,
            'note':'Try automatic retrieval or add a provider download link.'}


def download_one(row):
    import xray
    isin = row['isin']
    if row.get('override'):
        return source_meta(row)
    url = row.get('source_url') or row.get('provider_page') or ishares_page(isin)
    if not url:
        raise ValueError('Automatic source discovery is not available for this provider yet. Add its holdings download or product-page link.')
    try:
        data, final_url = fetch(url)
    except Exception:
        if isin not in ISHARES or row.get('source_url'):
            raise
        # Fallback: current UK workbook, checked against its Key Facts ISIN.
        download_url = ('https://www.blackrock.com/varnish-api/uk-retail01-product-data/'
            'product-data/api/v1/get-fund-document?' + urllib.parse.urlencode({
                'appType':'PRODUCT_PAGE', 'appSubType':'ISHARES',
                'targetSite':'ishares-uk', 'locale':'en_GB',
                'portfolioId':ISHARES[isin], 'component':'fundDownloadV2',
                'userType':'individual'}))
        data, final_url = fetch(download_url)
    verified = False
    asof = None
    if b'<html' in data[:1500].lower() or b'<!doctype html' in data[:1500].lower():
        page = html.unescape(data.decode('utf-8','replace'))
        host = urllib.parse.urlsplit(final_url).hostname or ''
        if not isin in page:
            raise ValueError('The product page does not contain this ISIN. Check the share class.')
        # Read an actual link, rather than fabricating a legacy CSV endpoint.
        links = re.findall(r'href=["\']([^"\']+)["\']',page)
        downloads = [urllib.parse.urljoin(final_url,l) for l in links
                     if 'get-fund-document?' in l or 'topholdingscsv' in l or
                     (('holding' in l.lower() or 'constituent' in l.lower()) and
                      re.search(r'\.csv|\.xlsx|fileType=csv',l,re.I))]
        if not downloads:
            raise ValueError('No supported full-holdings download found on this page. Paste the direct provider download link.')
        data, final_url = fetch(downloads[0])
        verified = True
    if b'urn:schemas-microsoft-com:office:spreadsheet' in data[:500]:
        data, asof = xml_holdings(data, isin)
        verified = True
    if b'<html' in data[:1500].lower() or b'<!doctype html' in data[:1500].lower():
        raise ValueError('The provider returned a web page, not holdings. Existing data has been kept.')
    DATA.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(dir=DATA) as td:
        p = Path(td)/('download.xlsx' if data[:2]==b'PK' else 'download.csv')
        p.write_bytes(data)
        frame = xray.read_holdings_file(p)
        if 'globalxetfs.eu' in urllib.parse.urlsplit(final_url).netloc:
            raw_rows = csv.DictReader(io.StringIO(data.decode('utf-8-sig')))
            dates = {r.get('AS_OF_DATE') for r in raw_rows} - {None, ''}
            if len(dates) == 1:
                asof = next(iter(dates))
        if data[:2] == b'PK' and not asof:
            from openpyxl import load_workbook
            wb = load_workbook(p, read_only=True, data_only=True)
            try:
                for cells in wb.worksheets[0].iter_rows(max_row=12, values_only=True):
                    for value in cells:
                        if isinstance(value, datetime):
                            asof = value.date().isoformat()
                        elif value:
                            match = re.search(r'(\d{1,2}[-/][A-Za-z]{3,4}[-/]\d{4})', str(value))
                            if match:
                                asof = match.group(1)
            finally:
                wb.close()
        coverage = float(frame.weight.sum())
        if frame.empty or not math.isfinite(coverage) or not 90 <= coverage <= 105:
            raise ValueError(f'Download rejected: holdings sum to {coverage:.2f}%. Expected a full holdings file (90–105%).')
        # Small cash liabilities are normal in long-only funds. Preserve their
        # signed weight, but reject short equity / leveraged exposure models.
        negative_equity = (frame.weight < 0) & (frame.bucket != 'Cash/Derivatives')
        if negative_equity.any() or float(frame.loc[frame.weight < 0, 'weight'].sum()) < -2:
            raise ValueError('Material short exposure needs a dedicated model; existing data kept.')
        # Persist a canonical CSV accepted by the existing engine.
        frame.rename(columns={'weight':'Weight (%)','aclass':'Asset Class'}).to_csv(Path(td)/'normalized.csv',index=False)
        if not asof:
            m = re.search(rb'(\d{1,2}/[A-Za-z]{3,4}/\d{4})',data[:1500])
            asof = m.group(1).decode() if m else None
        meta = {'source_url':final_url,'provider_page':url,'as_of':asof,'retrieved_at':now(),
                'holdings_count':len(frame),'coverage':coverage,'identity_verified':verified,
                'note': 'Share-class identity checked against provider.' if verified else
                        'Direct link supplied by you; fund identity is not embedded in this file. Check the source.'}
        with LOCK:
            (Path(td)/'normalized.csv').replace(snapshot_path(isin))
            atomic_json(DATA/(isin+'.json'),meta)
    return {**meta,'status':'downloaded','label':'Provider download'}


def bootstrap():
    draft = read_json(DATA/'portfolio.json',None)
    rows = draft['rows'] if draft else catalog()
    identities = read_json(DATA/'identities.json', {})
    rows = [{**r, **identities.get(r['isin']+':'+r['currency'], {})} for r in rows]
    return {'ok':True,'rows':[{**r,'source':source_meta(r)} for r in rows],
            'saved_at':draft.get('saved_at') if draft else None,
            'origin':'saved' if draft else 'manual', 'catalog':catalog()}


def start_refresh(rows):
    rows = validate_rows(rows)
    jid = uuid.uuid4().hex
    with LOCK:
        if any(not j['done'] for j in JOBS.values()):
            raise ValueError('A holdings refresh is already running. Wait for it to finish.')
        JOBS.clear()
        JOBS[jid] = {'done':False,'completed':0,'total':len(rows),'results':[]}
    def run():
        def one(r):
            try:
                return {'isin':r['isin'],'ticker':r['ticker'],'ok':True,'source':download_one(r)}
            except Exception as e:
                return {'isin':r['isin'],'ticker':r['ticker'],'ok':False,'message':str(e), 'source':source_meta(r)}
        with ThreadPoolExecutor(max_workers=3) as pool:
            for result in pool.map(one,rows):
                with LOCK:
                    JOBS[jid]['results'].append(result)
                    JOBS[jid]['completed'] += 1
        with LOCK:
            JOBS[jid]['done'] = True
    threading.Thread(target=run,daemon=True).start()
    return {'ok':True,'job':jid}


def start_verify(rows):
    rows = validate_rows(rows)
    jid = uuid.uuid4().hex
    with LOCK:
        if any(not j['done'] for j in JOBS.values()):
            raise ValueError('Another data check is running. Wait for it to finish.')
        JOBS.clear()
        JOBS[jid] = {'done':False, 'completed':0, 'total':len(rows), 'results':[]}
    def record(result):
        with LOCK:
            JOBS[jid]['results'].append(result)
            JOBS[jid]['completed'] += 1
    def run():
        def lookup(ib):
            from ib_async import Contract
            for row in rows:
                try:
                    details = ib.reqContractDetails(Contract(secType='STK', secIdType='ISIN',
                        secId=row['isin'], currency=row['currency'], exchange='SMART'))
                    records = [detail_record(d) for d in details]
                    exact = [r for r in records if r['isin']==row['isin'] and r['currency']==row['currency']]
                    if len(exact) != 1:
                        raise ValueError('Multiple or no matching listings; use investment search to choose explicitly.')
                    identity = {k:exact[0][k] for k in ('conid','exchange','identity_source')}
                    identity['verified_at'] = now()
                    with LOCK:
                        saved = read_json(DATA/'identities.json', {})
                        saved[row['isin']+':'+row['currency']] = identity
                        atomic_json(DATA/'identities.json', saved)
                    record({'isin':row['isin'],'ticker':row['ticker'],'ok':True,'identity':identity})
                except Exception as e:
                    record({'isin':row['isin'],'ticker':row['ticker'],'ok':False,'message':str(e)})
        try:
            with_broker(lookup)
        except Exception as e:
            with LOCK:
                done = {r['isin'] for r in JOBS[jid]['results']}
            for r in rows:
                if r['isin'] not in done:
                    record({'isin':r['isin'],'ticker':r['ticker'],'ok':False,
                            'message':'Paper Gateway lookup unavailable: '+type(e).__name__})
        finally:
            with LOCK: JOBS[jid]['done'] = True
    threading.Thread(target=run,daemon=True).start()
    return {'ok':True,'job':jid}


def analyze_rows(rows, auto_fetch=False):
    import pandas as pd
    import xray
    rows = validate_rows(rows,require_total=True)
    warnings, sources = [], []
    if auto_fetch:
        def missing(r):
            if source_meta(r)['status'] != 'missing':
                return None
            try:
                download_one(r)
            except Exception as e:
                return f"{r['ticker']}: automatic download unavailable: {e}"
        with ThreadPoolExecutor(max_workers=3) as pool:
            warnings.extend(w for w in pool.map(missing, rows) if w)
    with tempfile.TemporaryDirectory() as td:
        directory = Path(td)
        pf = []
        for row in rows:
            isin = row['isin']
            source = source_meta(row)
            sources.append({**row,'source':source})
            original = snapshot_path(isin)
            if not original.exists() and source.get('file'):
                original = HERE/'holdings'/source['file']
            if not row.get('override') and original.exists():
                # ISIN filenames prevent ticker collisions between venues/funds.
                shutil.copy(original,directory/(isin+original.suffix))
            if source['status']=='local':
                warnings.append(f"{row['ticker']}: local holdings, as of {source.get('as_of') or 'unknown date'}; not refreshed.")
            if source['status']=='missing':
                warnings.append(f"{row['ticker']}: no holdings. Its {row['weight']:g}% allocation is counted as unknown.")
            pf.append({**row,'ticker':isin,'url':''})
        enrich = HERE/'holdings'/'enrich.csv'
        if enrich.exists(): shutil.copy(enrich,directory/'enrich.csv')
        agg, exposures, summary, overlap = xray.analyze(pd.DataFrame(pf),directory)
    names = {r['isin']: (r['ticker'] if sum(x['ticker']==r['ticker'] for x in rows)==1
                        else r['ticker']+' · '+r['isin']) for r in rows}
    known, classified = 0,0
    for s in summary:
        if str(s[3]).startswith('override:'):
            classified += s[2]
        elif s[4]:
            known += s[2]*min(100,max(0,s[5]))/100
            if not 95 <= s[5] <= 105:
                warnings.append(f'{names[s[0]]}: parsed fund weights total {s[5]:.2f}%; interpret exposures with care.')
        else:
            warnings.append(f'{names[s[0]]}: holdings unavailable or could not be parsed; counted as unknown.')
    holdings=[]
    for key, h in agg.items():
        holdings.append({**h,'id':key,'etfs':{names.get(t,t):v for t,v in h['etfs'].items()}})
    holdings.sort(key=lambda h:-h['weight'])
    total = sum(h['weight'] for h in holdings)
    if abs(total-100)>.2:
        warnings.append(f'Provider rounding / exposure weights total {total:.2f}% rather than exactly 100%. Values have not been silently rescaled.')
    pairs = [{'a':names[a],'b':names[b],'weight':float(overlap.loc[a,b])}
             for i,a in enumerate(overlap.index) for b in overlap.index[i+1:]]
    pairs.sort(key=lambda p:-p['weight'])
    if any(h['weight'] < 0 for h in holdings):
        warnings.append('Small negative cash / derivative balances are retained as signed exposure; charts show positive bars only.')
    return {'ok':True,'analyzed_at':now(),'rows':sources,'holdings':holdings,
        'exposures':{k:sorted([{'name':n,'weight':float(w)} for n,w in v.items()],key=lambda x:-x['weight'])
                     for k,v in exposures.items()},
        'overlap':pairs,'warnings':list(dict.fromkeys(warnings)),
        'stats':{'funds':len(rows),'positions':sum(1 for h in holdings if h['bucket'] not in ('Unknown','Cash','Gold') and h['weight']>0),
                 'holdings_coverage':known,'classified':classified,'unknown':max(0,100-known-classified),'exposure_total':total}}


def broker_holdings():
    """Current positions by ISIN, read-only, for the current-vs-target column.

    Uses the same paper-only guarded connection as instrument lookup. It reads
    portfolio positions and nothing else: it never places, cancels or modifies
    an order, and never touches the saved draft or the investing policy."""
    def lookup(ib):
        from ib_async import Contract
        account = (ib.managedAccounts() or [''])[0]
        out, unpriced = {}, []
        for item in ib.portfolio(account):
            if not item.position:
                continue
            isin = None
            try:
                details = ib.reqContractDetails(Contract(conId=item.contract.conId))
                ids = {sid.value for d in details
                       if d.contract.conId == item.contract.conId
                       for sid in (d.secIdList or []) if sid.tag == 'ISIN'}
                if len(ids) == 1:
                    isin = ids.pop()
            except (TimeoutError, ConnectionError):
                pass
            if not isin:
                continue
            price = item.marketPrice if item.marketPrice and item.marketPrice > 0 else None
            eur = (price * item.position) if (price and item.contract.currency == 'EUR') else None
            if eur is None:
                unpriced.append(isin)
            out[isin] = {'shares': float(item.position), 'value_eur': eur,
                         'currency': item.contract.currency}
        return {'ok': True, 'holdings': out, 'unpriced': unpriced, 'read_at': now()}
    try:
        fresh = with_broker(lookup)
    except Exception as exc:                      # noqa: BLE001 - any broker failure
        # Gateway down, wrong account, timeout: fall back to the last good read
        # rather than showing nothing. Never silently: the caller gets stale=True
        # plus the timestamp it was actually read, and the UI must label it.
        cached = read_json(DATA / 'holdings.json', None)
        if not cached:
            raise
        return {**cached, 'ok': True, 'stale': True,
                'reason': str(exc) or 'Gateway could not be reached.'}
    atomic_json(DATA / 'holdings.json', fresh)
    return {**fresh, 'stale': False}


def api(action, p):
    if action=='bootstrap': return bootstrap()
    if action=='search': return search(p.get('query',''),p.get('online',False))
    if action=='resolve': return resolve(p)
    if action=='save':
        rows=validate_rows(p.get('rows'))
        saved=now()
        with LOCK: atomic_json(DATA/'portfolio.json',{'rows':rows,'saved_at':saved})
        return {'ok':True,'saved_at':saved}
    if action=='refresh': return start_refresh(p.get('rows'))
    if action=='verify': return start_verify(p.get('rows'))
    if action=='job':
        with LOCK:
            j=JOBS.get(p.get('job'))
            if j is None: raise ValueError('Refresh session expired. Start a new refresh.')
            return {'ok':True,**j}
    if action=='holdings': return broker_holdings()
    if action=='analyze': return analyze_rows(p.get('rows'), auto_fetch=True)
    raise ValueError('Unknown workspace action.')
