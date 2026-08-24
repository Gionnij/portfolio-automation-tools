#!/usr/bin/env python3
"""Portfolio X-Ray - look-through exposure analysis for an ETF portfolio.

Usage:
    python xray.py [portfolio.xlsx] [-o xray_report.xlsx] [--holdings-dir holdings]

Reads the portfolio sheet (columns: Ticker, Weight %, optional ISIN/Override/
Holdings URL), loads each ETF's holdings from <holdings-dir>/<TICKER>.csv or
.xlsx (auto-downloading from 'Holdings URL' when the file is missing), and
writes an Excel report with aggregated look-through exposures:
stocks, countries, sectors, currencies, asset classes, ETF overlap matrix.

Override column (skip look-through for cash/commodity ETPs):
    CASH:EUR    -> counts as asset class Cash, currency EUR
    GOLD:GOLD   -> counts as asset class Gold, currency bucket "Gold"

Requires: pandas, openpyxl   (pip install pandas openpyxl)
"""

import argparse
import csv
import re
import sys
import unicodedata
import urllib.request
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- column maps

SYN = {  # field -> header synonyms (checked in order; exact match wins first)
    "name":     ["name", "holding", "holding name", "security", "security name",
                 "company", "constituent", "issuer name", "bezeichnung",
                 "fund holding", "instrument name"],
    "weight":   ["weight (%)", "weighting (%)", "% of net assets", "weight %",
                 "weighting %", "weight", "weighting", "gewichtung", "peso",
                 "% weight", "percent of fund", "index weight",
                 "net assets (%)", "% of net asset value", "% nav",
                 "percent of net assets", "% of fund", "fund weight"],
    "isin":     ["isin"],
    "ticker":   ["issuer ticker", "ticker", "symbol"],
    "country":  ["location", "country of risk", "country", "land", "domicile"],
    "sector":   ["gics sector", "sector", "industry", "branche", "icb industry",
                 "industry classification"],
    "currency": ["market currency", "trading currency", "local currency",
                 "currency", "ccy", "waehrung", "währung"],
    "aclass":   ["asset class", "instrument type"],
}
WEIGHT_EXCLUDE = ("market value", "notional", "nominal", "price", "shares", "par")

COUNTRY_MAP = {
    "united states": "United States", "usa": "United States", "us": "United States",
    "united states of america": "United States", "vereinigte staaten": "United States",
    "united kingdom": "United Kingdom", "uk": "United Kingdom",
    "great britain": "United Kingdom",
    "korea": "South Korea", "korea (south)": "South Korea",
    "republic of korea": "South Korea", "south korea": "South Korea",
    "taiwan": "Taiwan", "taiwan, province of china": "Taiwan",
    "china": "China", "hong kong": "Hong Kong",
    "netherlands": "Netherlands", "the netherlands": "Netherlands",
    "russian federation": "Russia", "czech republic": "Czechia",
    "viet nam": "Vietnam", "cash and/or derivatives": "Cash/Derivatives",
}

ISIN_COUNTRY = {
    "US": "United States", "CA": "Canada", "GB": "United Kingdom",
    "IE": "Ireland", "DE": "Germany", "FR": "France", "NL": "Netherlands",
    "CH": "Switzerland", "SE": "Sweden", "DK": "Denmark", "FI": "Finland",
    "NO": "Norway", "IT": "Italy", "ES": "Spain", "BE": "Belgium",
    "AT": "Austria", "PT": "Portugal", "LU": "Luxembourg", "GR": "Greece",
    "JP": "Japan", "AU": "Australia", "HK": "Hong Kong", "SG": "Singapore",
    "NZ": "New Zealand", "CN": "China", "TW": "Taiwan", "KR": "South Korea",
    "IN": "India", "BR": "Brazil", "MX": "Mexico", "ZA": "South Africa",
    "SA": "Saudi Arabia", "ID": "Indonesia", "TH": "Thailand",
    "MY": "Malaysia", "AE": "United Arab Emirates", "QA": "Qatar",
    "KW": "Kuwait", "PL": "Poland", "TR": "Turkey", "PH": "Philippines",
    "CL": "Chile", "HU": "Hungary", "CZ": "Czechia", "PE": "Peru",
    "CO": "Colombia", "EG": "Egypt", "VN": "Vietnam", "AR": "Argentina",
    "RU": "Russia", "RO": "Romania", "IL": "Israel", "IS": "Iceland",
    "KY": "Cayman Islands", "BM": "Bermuda", "JE": "Jersey",
    "GG": "Guernsey", "IM": "Isle of Man", "PA": "Panama", "UY": "Uruguay",
}
COUNTRY_CCY = {
    "United States": "USD", "Canada": "CAD", "United Kingdom": "GBP",
    "Switzerland": "CHF", "Sweden": "SEK", "Denmark": "DKK", "Norway": "NOK",
    "Japan": "JPY", "Australia": "AUD", "Hong Kong": "HKD",
    "Singapore": "SGD", "New Zealand": "NZD", "China": "CNY",
    "Taiwan": "TWD", "South Korea": "KRW", "India": "INR", "Brazil": "BRL",
    "Mexico": "MXN", "South Africa": "ZAR", "Saudi Arabia": "SAR",
    "Indonesia": "IDR", "Thailand": "THB", "Malaysia": "MYR",
    "United Arab Emirates": "AED", "Qatar": "QAR", "Kuwait": "KWD",
    "Poland": "PLN", "Turkey": "TRY", "Philippines": "PHP", "Chile": "CLP",
    "Israel": "ILS", "Czechia": "CZK", "Hungary": "HUF", "Vietnam": "VND",
    "Egypt": "EGP", "Iceland": "ISK",
    "Cayman Islands": "USD", "Bermuda": "USD", "Jersey": "GBP",
    "Guernsey": "GBP", "Isle of Man": "GBP",
    **{c: "EUR" for c in ("Germany", "France", "Netherlands", "Italy",
       "Spain", "Ireland", "Belgium", "Austria", "Portugal", "Finland",
       "Luxembourg", "Greece")},
}

TICKER_SUFFIX_COUNTRY = {  # Bloomberg-style exchange codes
    "US": "United States", "UN": "United States", "UW": "United States",
    "UQ": "United States", "UR": "United States", "UA": "United States",
    "CN": "Canada", "CT": "Canada", "CV": "Canada",
    "LN": "United Kingdom", "FP": "France", "GY": "Germany", "GR": "Germany",
    "SW": "Switzerland", "VX": "Switzerland", "SE": "Switzerland",
    "IM": "Italy",
    "NA": "Netherlands", "BB": "Belgium", "SM": "Spain", "PL": "Portugal",
    "AV": "Austria", "SS": "Sweden", "DC": "Denmark", "NO": "Norway",
    "FH": "Finland", "ID": "Ireland", "PW": "Poland", "GA": "Greece",
    "JP": "Japan", "JT": "Japan", "AU": "Australia", "NZ": "New Zealand",
    "HK": "Hong Kong", "SP": "Singapore",
    "CH": "China", "C1": "China", "C2": "China", "CG": "China", "CS": "China",
    "TT": "Taiwan", "KS": "South Korea", "KQ": "South Korea",
    "IN": "India", "IS": "India", "IB": "India",
    "TB": "Thailand", "MK": "Malaysia", "IJ": "Indonesia",
    "PM": "Philippines", "BZ": "Brazil", "MM": "Mexico", "CI": "Chile",
    "SJ": "South Africa", "AB": "Saudi Arabia", "UH": "United Arab Emirates",
    "QD": "Qatar", "TI": "Turkey", "IT": "Israel",
    # Reuters-style fallbacks that don't collide with the above
    "L": "United Kingdom", "PA": "France", "DE": "Germany", "MI": "Italy",
    "AS": "Netherlands", "ST": "Sweden", "HE": "Finland", "T": "Japan",
    "TW": "Taiwan", "SZ": "China", "TO": "Canada", "AX": "Australia",
}


def ticker_country(tkr):
    t = str(tkr).strip().upper()
    if not t or t == "NAN":
        return ""
    m = re.search(r"[ .]([A-Z0-9]{1,3})$", t)
    if m:
        return TICKER_SUFFIX_COUNTRY.get(m.group(1), "")
    if re.fullmatch(r"[A-Z]{1,5}", t):       # bare ticker = US listing
        return "United States"
    return ""


REGIONS = {
    "North America": ["United States", "Canada"],
    "Europe (developed)": ["United Kingdom", "Germany", "France", "Switzerland",
        "Netherlands", "Sweden", "Italy", "Spain", "Denmark", "Finland",
        "Norway", "Belgium", "Austria", "Ireland", "Portugal", "Luxembourg",
        "Iceland", "Israel", "Monaco", "Liechtenstein", "Jersey",
        "Guernsey", "Isle of Man"],
    "Asia-Pacific (developed)": ["Japan", "Australia", "Hong Kong",
        "Singapore", "New Zealand", "Macau"],
    "Emerging Markets": ["China", "Taiwan", "India", "South Korea", "Brazil",
        "Saudi Arabia", "South Africa", "Mexico", "Indonesia", "Thailand",
        "Malaysia", "United Arab Emirates", "Qatar", "Kuwait", "Poland",
        "Turkey", "Philippines", "Chile", "Greece", "Hungary", "Czechia",
        "Peru", "Colombia", "Egypt", "Vietnam", "Argentina", "Russia",
        "Romania", "Nigeria", "Morocco", "Pakistan", "Bahrain", "Oman",
        "Jordan", "Sri Lanka", "Kazakhstan", "Kenya", "Bangladesh",
        "Mauritius", "Cyprus", "Panama", "Uruguay", "Cayman Islands",
        "Bermuda"],
}
COUNTRY_TO_REGION = {c: r for r, cs in REGIONS.items() for c in cs}


def region_of(country):
    if country in COUNTRY_TO_REGION:
        return COUNTRY_TO_REGION[country]
    if country in ("Cash/Derivatives",):
        return "Cash/Derivatives"
    if country.startswith("- "):            # override buckets: "- Cash -" etc.
        return country.strip("- ").strip()
    if country in ("Unknown", ""):
        return "Unknown"
    return "Other"


CASH_PAT = re.compile(
    r"cash|liquidit|money market|margin|fx forward|forward|future|swap|"
    r"collateral|repo|treasur.*bill|deposit|e-mini|payabl|receivabl|"
    r"(?:pound|yen|won|dollar|yuan|renminbi|franc|krona|euro)$", re.I)

STOP_TOKENS = {
    "INC", "CORP", "CORPORATION", "LTD", "LIMITED", "PLC", "SA", "NV", "AG",
    "SE", "SPA", "AB", "AS", "ASA", "OYJ", "CO", "COMPANY", "HOLDING",
    "HOLDINGS", "GROUP", "ADR", "GDR", "ORD", "SHS", "REGISTERED", "REG",
    "CLASS", "CL", "A", "B", "C", "N", "NON-VOTING", "PREF", "THE", "NEW",
}


def norm_h(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s.strip().lower())


def name_key(name):
    """Normalized merge key for a holding name (used when ISIN is missing)."""
    s = unicodedata.normalize("NFKD", str(name).upper())
    s = s.encode("ascii", "ignore").decode()
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    toks = [t for t in s.split() if t not in STOP_TOKENS]
    return " ".join(toks[:4]) or s.strip()


def to_num(x):
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace("%", "").replace(" ", "").strip()
    if not s or s in {"-", "--", "n/a", "N/A"}:
        return float("nan")
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else \
            s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return float("nan")


def map_columns(row):
    """Map field -> column index for a candidate header row."""
    hs = [norm_h(c) for c in row]
    out = {}
    for field, syns in SYN.items():
        # pass 1: exact
        for syn in syns:
            for i, h in enumerate(hs):
                if h == syn and i not in out.values():
                    out[field] = i
                    break
            if field in out:
                break
        if field in out:
            continue
        # pass 2: prefix/contains, with exclusions for weight
        for syn in syns:
            for i, h in enumerate(hs):
                if i in out.values() or not h:
                    continue
                if field == "weight" and any(x in h for x in WEIGHT_EXCLUDE):
                    continue
                if h.startswith(syn) or (len(syn) > 5 and syn in h):
                    out[field] = i
                    break
            if field in out:
                break
    return out


# ------------------------------------------------------------- file ingestion

def scan_for_header(rows):
    """Return (header_index, colmap) for the first row that maps name+weight."""
    for i, r in enumerate(rows[:100]):
        cm = map_columns(r)
        if "name" in cm and "weight" in cm:
            return i, cm
    return None, None


def _header_error(rows):
    sample = [" | ".join(str(c) for c in r if str(c).strip())
              for r in rows[:6] if any(str(c).strip() for c in r)]
    return ValueError("no header row with name+weight columns found; "
                      "file starts with:\n    " + "\n    ".join(sample[:4]))


def read_holdings_file(path, enrich=None):
    """Parse an issuer holdings file (csv/tsv/xlsx) -> normalized DataFrame."""
    data = Path(path).read_bytes()
    if data[:2] == b"PK":  # xlsx (even if named .csv)
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            sheets = pd.read_excel(path, header=None, dtype=object,
                                   sheet_name=None)
        idx = rows = cm = None
        for raw in sheets.values():
            srows = raw.astype(object).where(raw.notna(), "").values.tolist()
            i, c = scan_for_header(srows)
            if i is not None:
                idx, cm, rows = i, c, srows
                break
            if rows is None:
                rows = srows
        if idx is None:
            raise _header_error(rows or [])
        body = rows[idx + 1:]
    else:
        text = None
        for enc in ("utf-8-sig", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        best = None
        for delim in (",", ";", "\t", "|"):
            rows = list(csv.reader(text.splitlines(), delimiter=delim))
            idx, cm = scan_for_header(rows)
            if idx is not None and (best is None or len(cm) > len(best[3])):
                best = (rows, delim, idx, cm)
        if best is None:
            raise _header_error(list(csv.reader(text.splitlines()))[:6])
        rows, _, idx, cm = best
        body = rows[idx + 1:]

    # a repeated header row marks a second appended table (some issuer files
    # contain the holdings list twice) - parse only the first table
    for j, r in enumerate(body):
        cm2 = map_columns(r)
        if "name" in cm2 and "weight" in cm2:
            body = body[:j]
            break

    width = max(cm.values()) + 1
    recs = []
    for r in body:
        r = list(r) + [""] * (width - len(r))
        rec = {f: r[i] for f, i in cm.items()}
        recs.append(rec)
    df = pd.DataFrame(recs)
    df["weight"] = df["weight"].map(to_num)
    df["name"] = df.get("name", "").astype(str).str.strip()
    df = df[df["weight"].notna() & (df["name"] != "")]
    df = df[~df["name"].str.match(r"(?i)^(total|sum|gesamt)\b")]
    if df["weight"].sum() > 150:            # safety net: duplicated rows
        df = df.drop_duplicates(subset=["name", "weight"])
    if 0 < df["weight"].sum() < 2:          # fractions -> percent
        df["weight"] *= 100
    for col in ("isin", "country", "sector", "currency", "aclass", "ticker"):
        if col not in df.columns:
            df[col] = ""
        df[col] = df[col].astype(str).str.strip().replace({"nan": "", "-": ""})
    df["country"] = df["country"].map(
        lambda c: COUNTRY_MAP.get(norm_h(c), c.title() if c.islower() else c) or "Unknown")
    df["currency"] = df["currency"].str.upper().map(
        lambda c: c if re.fullmatch(r"[A-Z]{3}", c or "") else "Unknown")
    df["isin"] = df["isin"].str.upper().map(
        lambda s: s if re.fullmatch(r"[A-Z]{2}[A-Z0-9]{10}", s or "") else "")
    # fill missing country from ISIN prefix, missing currency from country
    m = (df["country"] == "Unknown") & (df["isin"] != "")
    df.loc[m, "country"] = df.loc[m, "isin"].str[:2].map(ISIN_COUNTRY).fillna("Unknown")
    m = df["country"] == "Unknown"          # then from ticker exchange suffix
    df.loc[m, "country"] = df.loc[m, "ticker"].map(
        lambda t: ticker_country(t) or "Unknown")
    m = (df["currency"] == "Unknown") & (df["country"] != "Unknown")
    df.loc[m, "currency"] = df.loc[m, "country"].map(COUNTRY_CCY).fillna("Unknown")
    # manual overrides from enrich.csv (key = exact ISIN or name substring)
    if enrich is not None:
        for _, e in enrich.iterrows():
            key = str(e.get("key", "")).strip()
            if not key:
                continue
            if re.fullmatch(r"[A-Z]{2}[A-Z0-9]{10}", key.upper()):
                mask = df["isin"] == key.upper()
            else:
                mask = df["name"].str.contains(key, case=False, regex=False)
            for fld in ("country", "currency", "sector"):
                v = str(e.get(fld, "")).strip()
                if v and v != "nan":
                    df.loc[mask, fld] = v.upper() if fld == "currency" else v
    is_cash = (df["aclass"].str.contains(CASH_PAT, na=False) |
               df["name"].str.contains(CASH_PAT, na=False))
    df["bucket"] = "Equity"
    df.loc[is_cash, "bucket"] = "Cash/Derivatives"
    df.loc[is_cash & (df["country"] == "Unknown"), "country"] = "Cash/Derivatives"
    return df.reset_index(drop=True)


def download(url, dest):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        payload = r.read()
    if payload[:2] == b"PK" and dest.suffix != ".xlsx":
        dest = dest.with_suffix(".xlsx")
    dest.write_bytes(payload)
    return dest


# ----------------------------------------------------------------- portfolio

def read_portfolio(path):
    df = pd.read_excel(path, sheet_name=0, dtype=object)
    ren = {}
    for c in df.columns:
        n = norm_h(c)
        if n.startswith("ticker") or n == "etf symbol":
            ren[c] = "ticker"
        elif n.startswith("weight") or "percentage" in n or n == "%":
            ren[c] = "weight"
        elif n.startswith("override"):
            ren[c] = "override"
        elif "url" in n:
            ren[c] = "url"
        elif n.startswith(("name", "description")):
            ren[c] = "name"
        elif n.startswith("isin"):
            ren[c] = "isin"
        elif n.startswith("cat"):
            ren[c] = "category"
    df = df.rename(columns=ren)
    for col in ("override", "url", "name", "isin", "category"):
        if col not in df.columns:
            df[col] = ""
    df = df[df.get("ticker").notna()]
    df["ticker"] = df["ticker"].astype(str).str.strip().str.upper()
    df = df[df["ticker"] != ""]
    df["weight"] = df["weight"].map(to_num)
    df = df[df["weight"].notna() & (df["weight"] > 0)]
    if df["weight"].sum() <= 1.5:            # stored as Excel percent fractions
        df["weight"] *= 100
    for col in ("override", "url", "name", "isin", "category"):
        df[col] = df[col].astype(str).str.strip().replace({"nan": ""})
    return df.reset_index(drop=True)


# ---------------------------------------------------------------- aggregation

def analyze(pf, holdings_dir):
    agg = {}                                  # key -> holding record
    per_etf_weights = {}                      # ticker -> {key: within-ETF %}
    summary, parsed = [], []
    holdings_dir.mkdir(exist_ok=True)
    epath = holdings_dir / "enrich.csv"
    enrich = pd.read_csv(epath, dtype=str).fillna("") if epath.exists() else None
    if enrich is not None:
        print(f"  using {len(enrich)} manual enrichment rules from {epath}")

    # -------- pass 1: locate and parse every holdings file
    for _, row in pf.iterrows():
        t = row["ticker"]
        override = row.get("override", "")
        override = "" if pd.isna(override) else str(override).strip()
        if override:
            parsed.append(("override", row, override, None))
            continue
        f = None
        for ext in (".csv", ".xlsx", ".xls", ".tsv"):
            cand = holdings_dir / (t + ext)
            if cand.exists():
                f = cand
                break
        if f is None and row["url"]:
            try:
                print(f"  downloading holdings for {t} ...")
                f = download(row["url"], holdings_dir / (t + ".csv"))
            except Exception as e:
                print(f"  !! download failed for {t}: {e}")
        if f is None:
            print(f"  !! no holdings data for {t} - counted as Unknown")
            parsed.append(("missing", row, None, None))
            continue
        try:
            parsed.append(("ok", row, read_holdings_file(f, enrich), f))
        except Exception as e:
            print(f"  !! parse failed for {f.name}: {e}")
            parsed.append(("error", row, None, f))

    # -------- crosswalk so the same stock matches across files that have
    # ISINs (Xtrackers, UBS) and files that don't (iShares, Global X)
    xwalk = {}
    for kind, row, h, f in parsed:
        if kind == "ok":
            for _, hr in h.iterrows():
                if hr["isin"]:
                    xwalk.setdefault(name_key(hr["name"]), hr["isin"])

    # -------- pass 2: aggregate on unified keys
    for kind, row, h, f in parsed:
        t, w = row["ticker"], row["weight"]
        if kind == "override":
            klass, _, ccy = h.partition(":")   # h carries the override string
            klass = klass.strip().title() or "Other"
            ccy = ccy.strip().upper() or "Unknown"
            agg["OVR:" + t] = dict(
                name=row["name"] or t, isin=row["isin"],
                country="- " + klass + " -", sector=klass,
                currency=ccy.title() if ccy == "GOLD" else ccy,
                bucket=klass, weight=w, etfs={t: w})
            summary.append((t, row["name"], w, "override: " + h, 1, 100.0, "-"))
            continue
        if kind in ("missing", "error"):
            src = "MISSING" if kind == "missing" else f"PARSE ERROR: {f.name}"
            agg["UNK:" + t] = dict(
                name=f"{row['name'] or t} (no holdings data)", isin="",
                country="Unknown", sector="Unknown", currency="Unknown",
                bucket="Unknown", weight=w, etfs={t: w})
            summary.append((t, row["name"], w, src, 0, 0.0, "-"))
            continue
        coverage = h["weight"].sum()
        if not 95 <= coverage <= 105:
            print(f"  !! {t}: parsed weights sum to {coverage:.1f}% - check {f.name}")
        scale = w / 100.0
        per_etf_weights[t] = {}
        for _, hr in h.iterrows():
            eff = hr["weight"] * scale
            nk = name_key(hr["name"])
            key = hr["isin"] or xwalk.get(nk) or nk
            per_etf_weights[t][key] = per_etf_weights[t].get(key, 0) + hr["weight"]
            rec = agg.setdefault(key, dict(name=hr["name"], isin=hr["isin"],
                                           country=hr["country"], sector=hr["sector"],
                                           currency=hr["currency"], bucket=hr["bucket"],
                                           weight=0.0, etfs={}))
            rec["weight"] += eff
            rec["etfs"][t] = rec["etfs"].get(t, 0) + eff
            if len(hr["name"]) > len(rec["name"]):
                rec["name"] = hr["name"]
            for fld in ("isin", "country", "sector", "currency"):
                if rec[fld] in ("", "Unknown") and hr[fld] not in ("", "Unknown"):
                    rec[fld] = hr[fld]
        if coverage < 100:                    # unreported remainder
            r = agg.setdefault("REM:" + t, dict(
                name=f"{t} unreported remainder", isin="", country="Unknown",
                sector="Unknown", currency="Unknown", bucket="Unknown",
                weight=0.0, etfs={}))
            r["weight"] += (100 - coverage) * scale
            r["etfs"][t] = (100 - coverage) * scale
        top = h.sort_values("weight", ascending=False).iloc[0]
        summary.append((t, row["name"], w, f.name, len(h), coverage,
                        f"{top['name']} ({top['weight']:.1f}%)"))

    # -------- exposures from final merged records, so fields backfilled from
    # any ETF (e.g. sectors for Global X names) apply everywhere
    exposures = {k: defaultdict(float) for k in
                 ("country", "sector", "currency", "aclass")}
    for rec in agg.values():
        exposures["country"][rec["country"] or "Unknown"] += rec["weight"]
        exposures["sector"][rec["sector"] or "Unknown"] += rec["weight"]
        exposures["currency"][rec["currency"] or "Unknown"] += rec["weight"]
        exposures["aclass"][rec["bucket"] or "Unknown"] += rec["weight"]
    exposures["region"] = defaultdict(float)
    for country, wv in exposures["country"].items():
        exposures["region"][region_of(country)] += wv

    # -------- pairwise overlap (sum of min within-ETF weights, %)
    tickers = list(per_etf_weights)
    overlap = pd.DataFrame(0.0, index=tickers, columns=tickers)
    for i, a in enumerate(tickers):
        for b in tickers[i:]:
            common = per_etf_weights[a].keys() & per_etf_weights[b].keys()
            v = sum(min(per_etf_weights[a][k], per_etf_weights[b][k]) for k in common)
            overlap.loc[a, b] = overlap.loc[b, a] = v
    return agg, exposures, summary, overlap


# -------------------------------------------------------------------- report

HDR_FILL = PatternFill("solid", start_color="2F5D50")
HDR_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
PCT = "0.00%"


def style_sheet(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def add_table(wb, title, header, rows, widths, pct_cols=(), total_col=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for r in rows:
        ws.append(list(r))
        for cell in ws[ws.max_row]:
            cell.font = BODY_FONT
    for c in pct_cols:
        for row in ws.iter_rows(min_row=2, min_col=c, max_col=c):
            row[0].number_format = PCT
    if total_col:
        n = ws.max_row
        L = get_column_letter(total_col)
        ws.cell(row=n + 1, column=total_col - 1, value="TOTAL").font = \
            Font(name="Arial", bold=True, size=10)
        tc = ws.cell(row=n + 1, column=total_col, value=f"=SUM({L}2:{L}{n})")
        tc.number_format, tc.font = PCT, Font(name="Arial", bold=True, size=10)
    style_sheet(ws, widths)
    return ws


def add_bar_chart(ws, title, items, anchor="G2", label_col=4, val_col=5):
    """Bar chart fed from a hidden helper range written in ASCENDING order,
    so every spreadsheet app renders the largest bar on top by default -
    no axis-reversal tricks, no renderer-dependent behavior."""
    ws.cell(row=1, column=val_col, value="Portfolio %")
    items = list(items)[::-1]                  # ascending
    for i, (k, v) in enumerate(items, 2):
        ws.cell(row=i, column=label_col, value=k)
        c = ws.cell(row=i, column=val_col, value=v / 100)
        c.number_format = PCT
    for col in (label_col, val_col):
        ws.column_dimensions[get_column_letter(col)].hidden = True
    n = len(items)
    ch = BarChart()
    ch.type, ch.title, ch.legend = "bar", title, None
    ch.height, ch.width = max(8, 0.5 * n + 2), 18
    data = Reference(ws, min_col=val_col, min_row=1, max_row=n + 1)
    cats = Reference(ws, min_col=label_col, min_row=2, max_row=n + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.y_axis.delete = False                   # show category labels
    ch.x_axis.delete = False
    ch.x_axis.numFmt = "0%"
    ch.gapWidth = 60
    ch.visible_cells_only = False              # helper columns are hidden
    ws.add_chart(ch, anchor)


def write_report(out, pf, agg, exposures, summary, overlap):
    wb = Workbook()
    wb.remove(wb.active)

    # Summary
    rows = [(t, n, w / 100, src, cnt, cov / 100, top)
            for t, n, w, src, cnt, cov, top in summary]
    ws = add_table(wb, "Summary",
                   ["Ticker", "ETF name", "Weight", "Holdings source",
                    "# holdings", "Coverage", "Largest position in ETF"],
                   rows, [10, 38, 9, 26, 11, 10, 42], pct_cols=(3, 6), total_col=3)

    # Holdings (look-through)
    recs = sorted(agg.values(), key=lambda r: -r["weight"])
    hrows = []
    for i, r in enumerate(recs, 1):
        via = "; ".join(f"{t} {v:.2f}%" for t, v in
                        sorted(r["etfs"].items(), key=lambda x: -x[1]))
        hrows.append((i, r["name"], r["isin"], r["country"], r["sector"],
                      r["currency"], r["weight"] / 100, via))
    add_table(wb, "Holdings",
              ["#", "Holding", "ISIN", "Country", "Sector", "Currency",
               "Portfolio %", "Via ETFs"],
              hrows, [5, 38, 15, 16, 22, 9, 11, 46], pct_cols=(7,), total_col=7)

    # Exposure sheets
    for sheet, key, label in (("Regions", "region", "Macro region"),
                              ("Countries", "country", "Country"),
                              ("Sectors", "sector", "Sector"),
                              ("Currencies", "currency", "Currency"),
                              ("Asset classes", "aclass", "Asset class")):
        items = sorted(exposures[key].items(), key=lambda x: -x[1])
        ws = add_table(wb, sheet, [label, "Portfolio %"],
                       [(k, v / 100) for k, v in items], [26, 12],
                       pct_cols=(2,), total_col=2)
        add_bar_chart(ws, sheet, items[:20], anchor="G2")

    # Overlap matrix
    if len(overlap) > 1:
        ws = wb.create_sheet("Overlap")
        ws.append(["Overlap %"] + list(overlap.columns))
        diag_fill = PatternFill("solid", start_color="D9D9D9")
        for r, t in enumerate(overlap.index, 2):
            ws.append([t] + [overlap.loc[t, c] / 100 for c in overlap.columns])
            for j, cell in enumerate(ws[ws.max_row][1:]):
                cell.number_format = PCT
                cell.font = BODY_FONT
                if j == r - 2:                # diagonal: always 100%, no info
                    cell.value = None
                    cell.fill = diag_fill
            ws[ws.max_row][0].font = Font(name="Arial", bold=True, size=10)
        n = len(overlap.columns)
        rng = f"B2:{get_column_letter(n + 1)}{n + 1}"
        ws.conditional_formatting.add(rng, ColorScaleRule(
            start_type="num", start_value=0, start_color="FFFFFF",
            mid_type="percentile", mid_value=60, mid_color="FBE5A2",
            end_type="max", end_color="C0504D"))
        style_sheet(ws, [12] + [9] * n)
        ws["A1"].value = "min-weight overlap"

    wb.save(out)


# ----------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description="Portfolio X-Ray")
    ap.add_argument("portfolio", nargs="?", default="portfolio.xlsx")
    ap.add_argument("-o", "--output", default="xray_report.xlsx")
    ap.add_argument("--holdings-dir", default="holdings")
    args = ap.parse_args()

    pf = read_portfolio(args.portfolio)
    print(f"Portfolio: {len(pf)} positions, total weight {pf['weight'].sum():.1f}%")
    if abs(pf["weight"].sum() - 100) > 0.01:
        print(f"  !! weights sum to {pf['weight'].sum():.1f}%, not 100%")

    agg, exposures, summary, overlap = analyze(pf, Path(args.holdings_dir))
    write_report(args.output, pf, agg, exposures, summary, overlap)

    recs = sorted(agg.values(), key=lambda r: -r["weight"])[:10]
    print("\nTop 10 look-through positions:")
    for r in recs:
        print(f"  {r['weight']:5.2f}%  {r['name']}")
    print(f"\nReport written to {args.output}")


if __name__ == "__main__":
    main()
